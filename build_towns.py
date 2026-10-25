"""Build data/towns.json from the Central Bureau of Statistics localities file, for the map and town-type breakdowns.

Usage: python build_towns.py bycode2024.xlsx [table1.xlsx table8.xlsx]
(download from https://www.cbs.gov.il/he/publications/doclib/2019/ishuvim/bycode2024.xlsx; needs openpyxl). The optional
tables are tables 1 and 8 of the bureau's socio-economic index 2021 (local authorities, and localities within regional
councils), which give each town its cluster from 1 (lowest) to 10.

Output, keyed by the town code the elections data uses: [English name, district, x, y, population group, kind, population,
Hebrew name, socio-economic cluster or 0], where x and y are Israeli grid (ITM) metres, so they can be drawn on a map as
they are. The Hebrew name keeps the hyphens, quote marks and brackets that the committee's files drop; app.py uses it in
their place."""
import json
import re
import sys
from pathlib import Path

import openpyxl

GROUPS = {1: 'Jewish', 2: 'Arab', 3: 'Bedouin', 4: 'Mixed'}


def kind(form, population):
    """A plain-language kind of place from the bureau's settlement-type code and population."""
    form = form or 0
    if form in (310, 320):
        return 'Moshav'
    if form == 330:
        return 'Kibbutz'
    if form in (350, 370):
        return 'Community settlement'
    if form in (450, 460):
        return 'Bedouin locality'
    if 100 <= form < 300:
        p = population or 0
        return 'Big city (100,000+)' if p >= 100_000 else 'City (20,000-100,000)' if p >= 20_000 else 'Small town (under 20,000)'
    if form == 440:
        return 'Small town (under 20,000)'
    return 'Other'


def clusters(table1, table8):
    """Town code -> socio-economic cluster 2021: a local authority's own (table 1: code in B, cluster in G), or a village's
    within its regional council (table 8: code in F, cluster in M)."""
    out = {}
    for path, code, cluster in ((table1, 1, 6), (table8, 5, 12)):
        for r in openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True):
            if len(r) > cluster and isinstance(r[code], (int, float)) and isinstance(r[cluster], (int, float)):
                out.setdefault(str(int(r[code])), int(r[cluster]))
    return out


def main(path, table1=None, table8=None):
    socio = clusters(table1, table8) if table8 else {}
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    col = {h: i for i, h in enumerate(rows[0])}
    pop_col = next(h for h in col if str(h).startswith('סך הכל אוכלוסייה'))
    towns = {}
    for r in rows[1:]:
        code, xy = r[col['סמל יישוב']], str(r[col['קואורדינטות']] or '')
        if not code or len(xy) != 12:
            continue
        population = r[col[pop_col]]
        towns[str(code)] = [r[col['שם יישוב באנגלית']] or '', r[col['שם מחוז']] or '', int(xy[:6]), int(xy[6:]),
                            GROUPS.get(r[col['דת יישוב']], 'Other'), kind(r[col['צורת יישוב שוטפת']], population), population or 0,
                            re.sub(r'\s*-\s*', '-', re.sub(r'\s+', ' ', str(r[col['שם יישוב']]).replace('*', ''))).strip(),
                            socio.get(str(code), 0)]
    out = Path(__file__).resolve().parent / 'data' / 'towns.json'
    out.parent.mkdir(exist_ok=True)
    out.write_text(json.dumps(towns, ensure_ascii=False, separators=(',', ':')), 'utf-8')
    print(f'{len(towns)} towns written to {out}, {sum(1 for t in towns.values() if t[8])} with a socio-economic cluster')


if __name__ == '__main__':
    main(*sys.argv[1:4])
